"""Flask aplikacija za praćenje vijesti (press clipping)."""
import io
import os
from collections import Counter
from datetime import datetime

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from flask import Flask, jsonify, render_template, request, send_file

import db
from config import KEYWORDS, KEYWORD_VARIATIONS
from fetcher import check_feeds
from scheduler import start_scheduler
from youtube_fetcher import check_youtube

app = Flask(__name__)

CHANNEL_LABELS = {"google_news": "Google News", "rss": "Direktni RSS", "youtube": "YouTube"}

# Sentinel vrijednost za opciju "Samo YouTube" u filteru "Izvor" - filtrira po
# kanalu (channel='youtube') umjesto po tačnom nazivu izvora/kanala.
YOUTUBE_SOURCE_FILTER = "__youtube__"

# Kanali koji se smatraju "Vijestima" u brzom filteru tipa sadržaja (pill
# dugmad iznad liste članaka) - sve osim YouTube-a.
NEWS_CHANNELS = ["rss", "google_news"]

SENTIMENT_LABELS = {"positive": "Pozitivno", "negative": "Negativno", "neutral": "Neutralno"}


def channel_label(channel):
    return CHANNEL_LABELS.get(channel, channel)


def sentiment_label(sentiment):
    return SENTIMENT_LABELS.get(sentiment, "Neutralno")


app.jinja_env.globals["channel_label"] = channel_label
app.jinja_env.globals["sentiment_label"] = sentiment_label
app.jinja_env.globals["YOUTUBE_SOURCE_FILTER"] = YOUTUBE_SOURCE_FILTER

# Sve ključne riječi i varijacije iz config.py, bez duplikata, uvijek
# prikazane u filteru bez obzira da li trenutno imaju članaka u bazi.
ALL_KEYWORDS = list(dict.fromkeys(KEYWORDS + KEYWORD_VARIATIONS))


def get_filters():
    return {
        "source": request.args.get("source", "").strip(),
        "keyword": request.args.get("keyword", "").strip(),
        "date_from": request.args.get("date_from", "").strip(),
        "date_to": request.args.get("date_to", "").strip(),
        "sentiment": request.args.get("sentiment", "").strip(),
        "content_type": request.args.get("content_type", "").strip(),
    }


def resolve_channel_filter(filters):
    """Kombinuje "Samo YouTube" opciju iz filtera Izvor i brzi filter tipa
    sadržaja (pill dugmad "Vijesti"/"YouTube") u channel/channel_in
    argumente za db.query_articles."""
    if filters["source"] == YOUTUBE_SOURCE_FILTER or filters["content_type"] == "youtube":
        return "youtube", None
    if filters["content_type"] == "news":
        return None, NEWS_CHANNELS
    return None, None


@app.route("/")
def index():
    filters = get_filters()
    is_youtube_only = filters["source"] == YOUTUBE_SOURCE_FILTER
    channel_value, channel_in_value = resolve_channel_filter(filters)
    articles = db.query_articles(
        source=None if is_youtube_only else (filters["source"] or None),
        channel=channel_value,
        channel_in=channel_in_value,
        keyword=filters["keyword"] or None,
        date_from=filters["date_from"] or None,
        date_to=filters["date_to"] or None,
        sentiment=filters["sentiment"] or None,
    )

    all_articles = db.query_articles()
    today = datetime.now().strftime("%Y-%m-%d")
    articles_today = sum(
        1 for a in all_articles if (a["published_date"] or "").startswith(today)
    )
    source_counts = Counter(a["source"] for a in all_articles)
    top_source, top_source_count = (
        source_counts.most_common(1)[0] if source_counts else ("—", 0)
    )

    return render_template(
        "index.html",
        articles=articles,
        sources=db.get_distinct_sources(),
        keywords=ALL_KEYWORDS,
        filters=filters,
        last_check=db.get_meta("last_check"),
        total=len(articles),
        total_all=len(all_articles),
        articles_today=articles_today,
        top_source=top_source,
        top_source_count=top_source_count,
        today=today,
    )


@app.route("/trends")
def trends():
    all_articles = db.query_articles()
    return render_template(
        "trends.html",
        total_all=len(all_articles),
        last_check=db.get_meta("last_check"),
    )


@app.route("/api/trends")
def api_trends():
    return jsonify(
        {
            "by_day": db.get_articles_by_day(days=30),
            "top_sources": db.get_top_sources(limit=10),
            "sentiment": db.get_sentiment_distribution(),
            "by_channel": db.get_channel_distribution(),
        }
    )


@app.route("/check-now", methods=["POST"])
def check_now():
    new_count = check_feeds()
    new_count += check_youtube()
    return jsonify(
        {
            "status": "ok",
            "new_articles": new_count,
            "last_check": db.get_meta("last_check"),
        }
    )


@app.route("/export")
def export_excel():
    filters = get_filters()
    is_youtube_only = filters["source"] == YOUTUBE_SOURCE_FILTER
    channel_value, channel_in_value = resolve_channel_filter(filters)
    articles = db.query_articles(
        source=None if is_youtube_only else (filters["source"] or None),
        channel=channel_value,
        channel_in=channel_in_value,
        keyword=filters["keyword"] or None,
        date_from=filters["date_from"] or None,
        date_to=filters["date_to"] or None,
        sentiment=filters["sentiment"] or None,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Press clipping"

    headers = ["Naslov", "Izvor", "Datum objave", "Ključna riječ", "Kanal", "Sentiment", "Isječak teksta", "Link"]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        ws.cell(row=1, column=col_idx).font = Font(bold=True)

    for article in articles:
        ws.append(
            [
                article["title"],
                article["source"],
                article["published_date"],
                article["keyword"],
                channel_label(article["channel"]),
                sentiment_label(article["sentiment"]),
                article["snippet"],
                article["link"],
            ]
        )

    widths = [45, 20, 20, 22, 16, 14, 55, 45]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"press_clipping_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    db.init_db()
    # Railway (i drugi PaaS) dodjeljuju port kroz PORT env varijablu i tada
    # želimo produkcijski način rada (bez debug reloadera). Lokalno, kad PORT
    # nije postavljen, ostajemo u debug modu kao i do sad.
    PORT = int(os.environ.get("PORT", 5000))
    DEBUG_MODE = os.environ.get("PORT") is None
    # Sa Flask reloaderom (debug=True) ovaj fajl se učita dva puta - scheduler
    # pokrećemo samo u procesu koji zaista opslužuje zahtjeve.
    if os.environ.get("WERKZEUG_RUN_MAIN") == "true" or not DEBUG_MODE:
        start_scheduler()
    app.run(debug=DEBUG_MODE, host="0.0.0.0", port=PORT)
